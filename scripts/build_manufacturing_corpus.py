"""Build the small manufacturing corpus with the bundled document Python runtime.

Never uploads, indexes, changes permissions, or calls a model. Refuses overwrites.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.manufacturing_corpus_data import (
    AS_OF, DOCUMENTS, EDITION, GROUPS, PUBLIC_REFERENCE, QUESTIONS, ROLE_PERMISSIONS,
)

# This edition uses one reviewed official download, not an arbitrary local PDF.
REFERENCE_SHA256 = "a798535274ef279adacce559217ae231b74c79ceadea7b1103bc7bcb15bf0acd"


def sha(raw):
    return hashlib.sha256(raw).hexdigest()


def dump(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def metadata_lines(item):
    rows = [f"文档编号：{item['document_code']}", f"版本：{item['business_version']}",
            f"生效日期：{item['effective_from']}", f"状态：{item['business_status']}",
            f"适用范围：{item['applicability']}"]
    if item.get("effective_to"):
        rows.append(f"失效日期：{item['effective_to']}")
    return rows


def make_docx(path, item):
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for style in doc.styles:
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2.0)
    for name, size in (("Normal", 10.5), ("Title", 17), ("Heading 1", 12), ("List Bullet", 10.5)):
        style = doc.styles[name]
        style.font.name = "SimSun"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "SimSun")
        style.paragraph_format.space_after = Pt(6)
    doc.add_paragraph(item["title"].replace("_", " ").replace("-", " "), "Title")
    for line in metadata_lines(item):
        doc.add_paragraph(line)
    for kind, value in item["sections"]:
        if kind == "heading":
            doc.add_heading(value, level=1)
        elif kind == "table":
            table = doc.add_table(rows=0, cols=len(value[0]))
            table.autofit = False
            widths = [2.5, 4.5, 2.3, 7.7] if len(value[0]) == 4 else [4.0] * len(value[0])
            for i, values in enumerate(value):
                cells = table.add_row().cells
                for cell, text, width in zip(cells, values, widths):
                    cell.width = Cm(width)
                    cell.text = text
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    properties = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement("w:tcBorders")
                    for edge in ("top", "left", "bottom", "right"):
                        border = OxmlElement(f"w:{edge}")
                        for key, attr in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                            border.set(qn(f"w:{key}"), attr)
                        borders.append(border)
                    properties.append(borders)
                    if i == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "E7EEF5")
                        properties.append(shading)
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
                if i == 0:
                    repeat = OxmlElement("w:tblHeader")
                    table.rows[0]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph()
        elif kind == "list":
            doc.add_paragraph(value, "List Bullet")
        else:
            doc.add_paragraph(value)
    header = section.header.paragraphs[0]
    header.text = "衡川精密  作业文件"
    header.style = doc.styles["Normal"]
    doc.core_properties.author = "制造业模拟语料维护者"
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 10, 28, tzinfo=timezone.utc)
    doc.save(path)


def pdf_bytes(item):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = "STSong-Light"
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(font))
    normal = ParagraphStyle("body", fontName=font, fontSize=10.5, leading=17, spaceAfter=7)
    heading = ParagraphStyle("heading", parent=normal, fontSize=12, leading=20, spaceBefore=8, keepWithNext=True)
    title = ParagraphStyle("title", parent=normal, fontSize=17, leading=25, spaceAfter=12)
    cell = ParagraphStyle("cell", parent=normal, fontSize=9, leading=14, alignment=TA_LEFT)
    stream = BytesIO()
    pdf = SimpleDocTemplate(stream, pagesize=(595, 842), leftMargin=38, rightMargin=38,
                            topMargin=40, bottomMargin=38, title=item["title"], author="制造业模拟语料维护者")
    story = [Paragraph(item["title"], title)]
    story += [Paragraph(line, normal) for line in metadata_lines(item)]
    story.append(Spacer(1, 8))
    for kind, value in item["sections"]:
        if kind == "page_break":
            story.append(PageBreak())
        elif kind == "table":
            widths = [75, 95, 75, 65, 209] if len(value[0]) == 5 else None
            rows = [[Paragraph(text, cell) for text in row] for row in value]
            table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#D9D9D9")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7EEF5")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]))
            story.extend([table, Spacer(1, 8)])
        else:
            style = heading if kind == "heading" else normal
            story.append(Paragraph(value.replace("\n", "<br/>"), style))

    def furniture(canvas, document):
        canvas.setFont(font, 8)
        canvas.drawString(38, 818, "衡川精密有限公司")
        canvas.drawRightString(557, 20, f"第{document.page}页")

    pdf.build(story, onFirstPage=furniture, onLaterPages=furniture)
    return stream.getvalue()


def make_scan(path, raw, qa_dir):
    from pdf2image import convert_from_bytes
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen.canvas import Canvas
    pages = convert_from_bytes(raw, dpi=200)
    qa_dir.mkdir(parents=True, exist_ok=True)
    canvas = Canvas(str(path), pagesize=(595, 842))
    for index, image in enumerate(pages, 1):
        preview = qa_dir / f"scan-source-{index}.png"
        image.save(preview)
        canvas.drawImage(ImageReader(image), 0, 0, width=595, height=842)
        canvas.showPage()
    canvas.save()


def make_text(path, item):
    markdown = item["format"] == "md"
    lines = [f"# {item['title']}" if markdown else item["title"], "", *metadata_lines(item), ""]
    for kind, value in item["sections"]:
        if kind == "heading":
            lines += [("## " if markdown else "") + value, ""]
        elif kind == "table":
            lines += ["| " + " | ".join(value[0]) + " |", "| " + " | ".join(["---"] * len(value[0])) + " |"]
            lines += ["| " + " | ".join(row) + " |" for row in value[1:]]
            lines.append("")
        elif kind == "code":
            lines += ["```json", value, "```", ""]
        elif kind == "list":
            lines += [("- " if markdown else "") + value, ""]
        else:
            lines += [value, ""]
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8-sig" if item["id"] == "MFG-17" else "utf-8")


def build(output, dataset, *, reference_file=None, download_public=False):
    if any(directory.exists() and any(directory.iterdir()) for directory in (output, dataset)):
        raise ValueError("输出目录必须为空；已有语料不会被覆盖，请选择新目录，验证已有文件请使用 validate_manufacturing_corpus.py")
    if reference_file is not None:
        reference = reference_file.read_bytes()
    elif download_public:
        from urllib.request import Request, urlopen
        request = Request(PUBLIC_REFERENCE["source_url"], headers={
            "User-Agent": "EnterpriseKnowledgeCorpus/1.0 (+https://github.com/Heresy0/Multi-tenant-long-context-rag-engine)"})
        with urlopen(request, timeout=45) as response:
            reference = response.read(20_000_001)
    else:
        raise ValueError("请提供 --reference-file，或显式指定 --download-public 从已登记的NIST官方地址获取")
    if not reference.startswith(b"%PDF-") or not 10_000 < len(reference) < 20_000_000:
        raise ValueError("公开参考文件不是预期大小的PDF")
    if sha(reference) != REFERENCE_SHA256:
        raise ValueError("公开参考PDF与本版核验的NIST原件哈希不一致，须复核来源并发布新语料版本")
    output.mkdir(parents=True, exist_ok=True)
    dataset.mkdir(parents=True, exist_ok=True)
    entries = []
    for item in DOCUMENTS:
        path = output / item["file"]
        path.parent.mkdir(parents=True, exist_ok=True)
        if item["format"] == "docx":
            make_docx(path, item)
        elif item["format"] == "pdf":
            raw = pdf_bytes(item)
            if item["scan"]:
                make_scan(path, raw, ROOT / "tmp/manufacturing-layout")
            else:
                path.write_bytes(raw)
        else:
            make_text(path, item)
        entries.append({**{key: value for key, value in item.items() if key != "sections"},
                        "sha256": sha(path.read_bytes()), "bytes": path.stat().st_size})
    path = output / PUBLIC_REFERENCE["file"]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(reference)
    entries.append({**PUBLIC_REFERENCE, "sha256": sha(reference), "bytes": len(reference),
                    "retrieved_at": datetime.now(timezone.utc).isoformat()})
    entries.sort(key=lambda row: row["id"])
    source = Path(__file__).with_name("manufacturing_corpus_data.py")
    dump(output / "manifest.json", dict(edition=EDITION, as_of=AS_OF, synthetic_documents=len(DOCUMENTS),
         public_documents=1, total_documents=len(entries), counts=dict(Counter(row["format"] for row in entries)),
         group_counts=dict(Counter(row["group"] for row in entries)), source_definition_sha256=sha(source.read_bytes()),
         notice="企业、产品、工艺数值及事件为合成设定；NIST原文单独标记，不代表企业批准要求。", documents=entries))
    rows = "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in QUESTIONS)
    (dataset / "single_turn.jsonl").write_text(rows, encoding="utf-8")
    dump(dataset / "role_permissions.json", dict(kind="offline_role_fixture_not_live_grants", roles=ROLE_PERMISSIONS))
    scan = next(item for item in DOCUMENTS if item["scan"])
    dump(output / "validation/scan_ground_truth.json", dict(document=Path(scan["file"]).name,
         purpose="独立对照标注，不供OCR、生成或检索读取", source_sha256=next(row["sha256"] for row in entries if row["id"] == scan["id"]),
         expected_fields=["WO-017", "HC-A17", "OP-20", "2026年10月26日14:10", "3件", "质量人员签署"],
         text="\n".join(value for kind, value in scan["sections"] if kind != "page_break")))
    return dict(documents=len(entries), questions=len(QUESTIONS), formats=dict(Counter(row["format"] for row in entries)), model_calls=0)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "sample_docs/manufacturing_lite")
    parser.add_argument("--dataset-dir", type=Path, default=ROOT / "evals/manufacturing_lite")
    parser.add_argument("--reference-file", type=Path)
    parser.add_argument("--download-public", action="store_true")
    args = parser.parse_args(argv)
    print(json.dumps(build(args.output_dir, args.dataset_dir, reference_file=args.reference_file,
                           download_public=args.download_public), ensure_ascii=False))


if __name__ == "__main__":
    main()
