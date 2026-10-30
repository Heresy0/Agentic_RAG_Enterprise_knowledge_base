"""Conservative cleaning boundaries, provenance, and adapter integration."""
import copy
import hashlib
import json
import sys

import pytest
from docx import Document as WordDocument

from backend.app.documents.chunk_quality import CLEANING_VERSION, summarize_quality
from backend.app.documents.document_cleaning import (
    clean_document, clean_table, clean_text, cleaning_trace, summarize_cleaning,
)
from backend.app.documents.document_splitter import ContentUnit
from backend.app.documents import pdf_splitter as pdf
from backend.app.knowledge.rag import split_file


@pytest.mark.parametrize('profile', ['pdf_text', 'text_file', 'prose', 'literal'])
def test_no_guessing_numbers_units_negation_names_or_versions(profile):
    text = '限额 -10 MPa;误差 0.15 mm;版本 1.4\n不得自动重试;负责人周未;批次 O01/l1。'
    cleaned, audit = clean_text(text, profile=profile)
    assert cleaned == text
    assert audit['changed'] is False
    assert audit['applied_rules'] == []


def test_file_normalization_preserves_code_spacing_and_compatibility_characters():
    text = '\ufeff# 规范\r\n```python\r\nif ok:\r\n\tprint("Ａ  B")\r\n    save()\r\n```\r\n\x00'
    cleaned, audit = clean_text(text, profile='text_file')
    assert cleaned == '# 规范\n```python\nif ok:\n\tprint("Ａ  B")\n    save()\n```\n'
    assert audit['applied_rules'] == ['remove_utf8_bom', 'normalize_newlines', 'remove_nul']


def test_table_cleaning_keeps_positions_and_values_without_mutating_source():
    cells = [['项目', '限额', '备注'], ['垫片', '0.15 mm', ''],
             ['紧固件', None, ' 不得  重试\n版本 1.4 '], ['', '', '']]
    original = copy.deepcopy(cells)
    rows, audit = clean_table(cells)
    assert rows == ['项目|限额|备注', '垫片|0.15 mm|', '紧固件||不得 重试 版本 1.4']
    assert cells == original
    assert audit['counts'] == dict(changed_cells=1, normalized_empty_cells=1, dropped_empty_rows=1)
    assert audit['input_sha256'] != audit['output_sha256']
    _, unchanged = clean_table([['项目', '限额'], ['垫片', '0.15 mm']])
    assert unchanged['changed'] is False
    assert unchanged['applied_rules'] == []


def test_common_entry_preserves_literal_blocks_structure_and_adapter_metadata():
    units = [
        ContentUnit('code', '```python\nif ok:\n\tprint("Ａ  B")\n    save()\n```', ('参数',)),
        ContentUnit('table', [r'部件\|A|`a|b`||0.15 mm', 'B||不得重试|1.4'], ('参数',)),
        ContentUnit('text', '  不得重试。\n批次 O01/l1。  ', ('约束',), {'source_section_id':'s-2'}),
    ]
    original = copy.deepcopy(units)
    cleaned = clean_document(units, source_format='md')
    assert units == original
    assert cleaned[0].content == original[0].content
    assert cleaned[1].content == original[1].content
    assert cleaned[1].content is not units[1].content
    assert cleaned[2].content == '不得重试。\n批次 O01/l1。'
    assert cleaned[2].metadata['source_section_id'] == 's-2'
    assert [u.section_path for u in cleaned] == [u.section_path for u in units]
    assert [u.content for u in clean_document(cleaned, source_format='md')] == [u.content for u in cleaned]


def test_audit_is_deterministic_content_free_and_hashes_the_observed_text():
    before = '  内部正文 TOKEN-92745  '
    after, audit = clean_text(before, profile='prose', scope={'unit':3})
    assert audit == clean_text(before, profile='prose', scope={'unit':3})[1]
    assert audit['input_sha256'] == hashlib.sha256(before.encode()).hexdigest()
    assert audit['output_sha256'] == hashlib.sha256(after.encode()).hexdigest()
    serialized = json.dumps(audit, ensure_ascii=False)
    assert 'TOKEN-92745' not in serialized
    assert '内部正文' not in serialized
    assert audit['version'] == CLEANING_VERSION


def test_pdf_page_traces_follow_only_applicable_source_pages():
    traces = [cleaning_trace(' x ', 'x', profile='pdf_text', scope={'page':i}) for i in (1, 2, 3)]
    unit = pdf.PdfUnit('paragraph', '规则正文。', ('规则',), 1, 2, ('pypdf',))
    cleaned = clean_document([unit], source_format='pdf', upstream_traces=traces)[0]
    assert (cleaned.page_start, cleaned.page_end, cleaned.parsers) == (1, 2, ('pypdf',))
    pages = {t['scope']['page'] for t in cleaned.metadata['cleaning']['traces'] if 'page' in t['scope']}
    assert pages == {1, 2}
    assert not unit.metadata


def test_reviewed_ocr_trace_does_not_clear_review_required(tmp_path, monkeypatch):
    path = tmp_path/'扫描件.pdf'
    path.write_bytes(b'source-bound OCR fixture')
    monkeypatch.setattr(pdf, 'extract_pdf_pages', lambda *a, **k: [pdf.ExtractedPage(1, '负责人周未。', 'ocr')])
    monkeypatch.setattr(pdf, 'apply_reviewed_ocr', lambda *a, **k: ('负责人周末。', [{'before':'周未', 'after':'周末'}]))
    chunks = pdf.split_pdf(path)
    assert '负责人周末' in chunks[0].page_content
    assert chunks[0].metadata['ocr_corrections'] == [{'before':'周未', 'after':'周末'}]
    assert all(c.metadata['quality']['review_required'] for c in chunks)
    assert summarize_cleaning(chunks)['applied_rules']['hash_bound_reviewed_ocr'] == 1


@pytest.mark.parametrize('suffix', ['.md', '.txt'])
def test_file_traces_survive_multiple_child_chunks_and_summary_deduplicates(tmp_path, suffix):
    path = tmp_path/f'规范{suffix}'
    path.write_bytes(('\ufeff规则必须保留。\r\n\r\n' + '不得自动重试，必须保留原记录。'*150).encode('utf-8'))
    chunks = split_file(path)
    assert len(chunks) > 1
    assert all(c.metadata['cleaning']['version'] == CLEANING_VERSION for c in chunks)
    assert all('_cleaning_traces' not in c.metadata for c in chunks)
    summary = summarize_cleaning(chunks)
    assert summary == summarize_cleaning(chunks*2)
    assert summary['applied_rules']['remove_utf8_bom'] == 1
    assert summary['applied_rules']['normalize_newlines'] == 1
    assert summarize_quality(chunks)['cleaning'] == summary


def test_docx_table_adapter_records_cell_cleaning_and_keeps_section(tmp_path):
    path = tmp_path/'工艺.docx'
    document = WordDocument()
    document.add_heading('工艺参数', level=1)
    table = document.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ['部件', '公差', '备注']):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, [' 垫片 ', '0.15 mm', '']):
        cell.text = value
    document.save(path)
    chunks = split_file(path)
    table_chunk = next(c for c in chunks if c.metadata['chunk_type'] == 'table')
    assert '垫片|0.15 mm|' in table_chunk.page_content
    assert table_chunk.metadata['section_path'] == '工艺参数'
    assert summarize_cleaning(chunks)['applied_rules']['normalize_cell_whitespace'] == 1


def test_inspection_defaults_to_hashes_without_duplicating_body(tmp_path, monkeypatch, capsys):
    from scripts.inspect_document_chunks import main
    path = tmp_path/'规范.md'
    path.write_text('# 操作要求\n\n内部正文 TOKEN-92745，不得重试。', encoding='utf-8')
    monkeypatch.setattr(sys, 'argv', ['inspect_document_chunks.py', str(path)])
    assert main() == 0
    output = capsys.readouterr().out
    result = json.loads(output)
    assert result['chunks'][0]['cleaning']['traces']
    assert 'content' not in result['chunks'][0]
    assert 'TOKEN-92745' not in output
